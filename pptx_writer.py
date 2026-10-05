"""pptx_writer.py - draws the report deck (PowerPoint) from a finished slide plan.

It does no analysis and never calls the AI. It receives numbers and sentences that
report_engine.py has already computed and checked, and turns them into slides with
native, editable charts and tables, in the visual style of the Diageo status deck:
soft pastel chart colours, green tables and rounded insight boxes.

Input (all keys optional unless marked):
    deck = {
        "title": "Monthly Report",            # small header on every slide + cover title
        "subtitle": "July 2026",              # cover subtitle
        "sources": ["clients.xlsx"],          # cover footnote
        "slides": [ slide, slide, ... ],
    }
Every content slide has: "title" (required), "insights": [str], "recommendation": str | [str],
and optionally "source": "file.xlsx :: Sheet" (small footnote).
Slide types:
    {"type": "bar" | "line" | "donut", "chart_title": str, "categories": [...],
     "series": [{"name": str, "values": [...]}], "stacked": bool, "horizontal": bool | None,
     "number_format": "#,##0"}
    {"type": "table", "columns": [...], "rows": [[...], ...]}
    {"type": "kpis", "kpis": [{"label": str, "value": str, "note": str}]}
    {"type": "summary", "title": str, "items": [{"heading": str, "text": str}]}
"""

from __future__ import annotations

import copy
import io
import math
import re
from pathlib import Path
from typing import Any

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

# ---------------------------------------------------------------------------
# Look (taken from the Diageo status deck)
# ---------------------------------------------------------------------------
FONT = "Century Gothic"
SLIDE_W, SLIDE_H = 13.333, 7.5
TEMPLATE_PATH = Path(__file__).parent / "templates" / "diageo_template.pptx"

TEXT = "262626"
GREY = "5E5E5E"
LIGHT_GREY = "A3A3A3"
GREEN = "6E9B72"  # table header and insight-box border
GREEN_TINT = "EAF5E9"
BLUE = "73AEF7"
BLUE_TINT = "EAF2FD"
COVER_BLUE = "1876F1"
PALETTE = ["73AEF7", "F49591", "8CC48F", "FFD966", "F7A6CC", "A3A3A3", "9DE0D0", "C9B6E4"]
LINE_COLORS = ["4A90E2", "F26B6B", "5FA864", "E0A800", "D6559B", "7F7F7F", "3FB8A0", "9B7FD1"]
TINTS = ["E3EEFD", "FDE6E4", "E3F1E2", "FFF4CC", "FCE4F0", "EEEEEE", "DFF5EF", "EEE7F8"]

MAX_CATEGORIES = 30
MAX_SERIES = 8
MAX_TABLE_ROWS = 12
MAX_KPIS = 8

_CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")
_NUMERIC_TEXT = re.compile(r"^[\s$€£₹+\-−(]*[\d.,]+\s*(%|[KMBkmb])?\)?\s*$")


# ---------------------------------------------------------------------------
# Small helpers
# ---------------------------------------------------------------------------
def _clean(value: Any) -> str:
    """Text safe for PowerPoint (control characters make python-pptx fail)."""
    return _CONTROL_CHARS.sub("", "" if value is None else str(value)).strip()


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def _number(value: Any) -> float | None:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return None if math.isnan(number) or math.isinf(number) else number


def _as_list(value: Any) -> list[str]:
    if value is None:
        return []
    if isinstance(value, str):
        value = [value]
    return [text for text in (_clean(item) for item in value) if text]


def _style_run(run, size: float, bold: bool = False, color: str = TEXT) -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = _rgb(color)


def _text(slide, x, y, w, h, text, size=12, bold=False, color=TEXT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = anchor
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = _clean(text)
    _style_run(run, size, bold, color)
    return box


def _box(slide, x, y, w, h, fill: str, line: str, radius: float = 0.04):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.adjustments[0] = radius
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(fill)
    shape.line.color.rgb = _rgb(line)
    shape.line.width = Pt(1)
    shape.shadow.inherit = False
    for effect in shape._element.xpath("./p:style/a:effectRef"):
        effect.set("idx", "0")  # no theme effects (shadow) in any viewer
    return shape


def _bullet(paragraph, indent_in: float = 0.18) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(indent_in))))
    pPr.set("indent", str(-int(Inches(indent_in))))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for element in pPr.findall(qn(tag)):
            pPr.remove(element)
    etree.SubElement(pPr, qn("a:buChar")).set("char", "•")


# ---------------------------------------------------------------------------
# Text fitting (Century Gothic is wide: about 0.56 em per character)
# ---------------------------------------------------------------------------
def _count_lines(text: str, chars_per_line: int) -> int:
    lines, current = 1, 0
    for word in text.split():
        need = len(word) + (1 if current else 0)
        if current + need > chars_per_line:
            lines += 1
            current = len(word)
        else:
            current += need
    return lines


def _height_needed(items: list[str], size: float, width_in: float, heading: bool) -> float:
    usable = width_in - 0.4 - 0.22  # box margins + bullet indent
    chars_per_line = max(8, int(usable / (size * 0.00778)))
    total = 0.36 if heading else 0.0
    for item in items:
        total += _count_lines(item, chars_per_line) * size * 1.22 / 72 + 0.07
    return total + 0.2


def _fit(items: list[str], width_in: float, height_in: float, heading: bool = True, sizes=(13, 12, 11, 10, 9)):
    """Largest font size that fits. If nothing fits, drop trailing items. Returns (size, items)."""
    for size in sizes:
        if _height_needed(items, size, width_in, heading) <= height_in:
            return size, items
    keep = list(items)
    while len(keep) > 1 and _height_needed(keep, sizes[-1], width_in, heading) > height_in:
        keep.pop()
    return sizes[-1], keep


def _list_box(slide, x, y, w, h, heading: str, items: list[str], fill: str, line: str) -> None:
    items = [i for i in items if i]
    shape = _box(slide, x, y, w, h, fill, line)
    frame = shape.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.2)
    frame.margin_top = frame.margin_bottom = Inches(0.12)
    frame.vertical_anchor = MSO_ANCHOR.TOP
    size, items = _fit(items, w, h)
    bullets = len(items) > 1

    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.LEFT
    paragraph.space_after = Pt(5)
    run = paragraph.add_run()
    run.text = heading
    _style_run(run, size + 1, True, TEXT)
    for item in items:
        paragraph = frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.space_after = Pt(4)
        run = paragraph.add_run()
        run.text = item
        _style_run(run, size, False, TEXT)
        if bullets:
            _bullet(paragraph)


# ---------------------------------------------------------------------------
# Slide frame
# ---------------------------------------------------------------------------
def _layout(prs, name: str):
    for layout in prs.slide_layouts:
        if layout.name == name:
            return layout
    return None


def _new_slide(prs):
    layout = _layout(prs, "Content")
    if layout is None:  # no template file: plain blank slide
        return prs.slides.add_slide(prs.slide_layouts[6])
    slide = prs.slides.add_slide(layout)
    for placeholder in layout.placeholders:  # python-pptx does not copy the footer and page number
        if placeholder.placeholder_format.type in (PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER):
            slide.shapes._spTree.append(copy.deepcopy(placeholder._element))
    return slide


def _drop_last_slide(prs) -> None:
    id_list = prs.slides._sldIdLst
    last = id_list[-1]
    prs.part.drop_rel(last.rId)
    id_list.remove(last)


def _frame(slide, deck_title: str, title: str, number: int, source: str | None) -> None:
    _text(slide, 0.5, 0.26, 9.0, 0.2, deck_title.upper(), size=9, color=LIGHT_GREY)
    title = _clean(title)[:110] or "Untitled"
    size = 22 if len(title) <= 62 else 18 if len(title) <= 85 else 16
    _text(slide, 0.5, 0.5, 12.3, 0.6, title, size=size, color=GREY)
    on_template = slide.slide_layout.name == "Content"  # the template draws the page number
    if _clean(source):                                                                                   
        x, w = 0.5, 10.8                                                                                  
        _text(slide, x, 7.05, w, 0.25, f"Source: {_clean(source)[:90]}", size=9, color=LIGHT_GREY)
    if not on_template:
        _text(slide, 12.0, 7.05, 0.83, 0.25, str(number), size=9, color=LIGHT_GREY, align=PP_ALIGN.RIGHT)


def _insight_boxes_right(slide, spec) -> None:
    """Layout A: chart on the left, insight and recommendation boxes stacked on the right."""
    _list_box(slide, 8.45, 1.25, 4.38, 3.55, "Business insights", _as_list(spec.get("insights")), GREEN_TINT, GREEN)
    _list_box(slide, 8.45, 4.95, 4.38, 1.95, "Recommendation(Experimental Feature)", _as_list(spec.get("recommendation")), BLUE_TINT, BLUE)


def _insight_boxes_bottom(slide, spec) -> None:
    """Layout B: content on top, insight and recommendation boxes side by side underneath."""
    _list_box(slide, 0.5, 5.2, 7.4, 1.7, "Business insights", _as_list(spec.get("insights")), GREEN_TINT, GREEN)
    _list_box(slide, 8.1, 5.2, 4.73, 1.7, "Recommendation(Experimental Feature)", _as_list(spec.get("recommendation")), BLUE_TINT, BLUE)


# ---------------------------------------------------------------------------
# Charts
# ---------------------------------------------------------------------------
def _add_chart(slide, spec: dict[str, Any], x: float, y: float, w: float, h: float) -> None:
    kind = spec.get("type")
    categories = [_clean(c) or " " for c in spec.get("categories", [])][:MAX_CATEGORIES]
    series_specs = list(spec.get("series", []))[:MAX_SERIES]
    if not categories or not series_specs:
        raise ValueError("the chart has no data")
    number_format = spec.get("number_format") or "#,##0"

    data = CategoryChartData(number_format=number_format)
    data.categories = categories
    for series in series_specs:
        values = [_number(v) for v in list(series.get("values", []))[: len(categories)]]
        values += [None] * (len(categories) - len(values))
        data.add_series(_clean(series.get("name")) or "Value", values)
    if kind != "donut" and all(v is None for s in series_specs for v in s.get("values", [])):
        raise ValueError("the chart has no numeric values")

    stacked = bool(spec.get("stacked")) and len(series_specs) > 1
    horizontal = False
    if kind == "donut":
        chart_type = XL_CHART_TYPE.DOUGHNUT
    elif kind == "line":
        chart_type = XL_CHART_TYPE.LINE_MARKERS
    else:
        horizontal = spec.get("horizontal")
        if horizontal is None:
            horizontal = max(len(c) for c in categories) > 14 and len(categories) <= 15
        horizontal = bool(horizontal)
        chart_type = {
            (False, False): XL_CHART_TYPE.COLUMN_CLUSTERED,
            (False, True): XL_CHART_TYPE.COLUMN_STACKED,
            (True, False): XL_CHART_TYPE.BAR_CLUSTERED,
            (True, True): XL_CHART_TYPE.BAR_STACKED,
        }[(horizontal, stacked)]

    chart = slide.shapes.add_chart(chart_type, Inches(x), Inches(y), Inches(w), Inches(h), data).chart
    chart.font.name = FONT
    chart.font.size = Pt(10)
    chart.font.color.rgb = _rgb(GREY)
    chart.has_title = False

    if kind == "donut" or len(series_specs) > 1:
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.RIGHT if kind == "donut" else XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(10)
        chart.legend.font.name = FONT
    else:
        chart.has_legend = False

    plot = chart.plots[0]
    point_count = len(categories) * len(series_specs)

    if kind == "donut":
        for index in range(len(categories)):
            point = plot.series[0].points[index]
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = _rgb(PALETTE[index % len(PALETTE)])
            point.format.line.color.rgb = _rgb("FFFFFF")
        for hole in chart._chartSpace.xpath(".//c:holeSize"):
            hole.set("val", "58")
    elif kind == "line":
        for index, series in enumerate(plot.series):
            color = _rgb(LINE_COLORS[index % len(LINE_COLORS)])
            series.smooth = False
            series.format.line.color.rgb = color
            series.format.line.width = Pt(2.5)
            series.marker.style = XL_MARKER_STYLE.CIRCLE
            series.marker.size = 7
            series.marker.format.fill.solid()
            series.marker.format.fill.fore_color.rgb = color
            series.marker.format.line.color.rgb = color
    else:
        plot.gap_width = 60
        if stacked:
            plot.overlap = 100
        for index, series in enumerate(plot.series):
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = _rgb(PALETTE[index % len(PALETTE)])
            series.invert_if_negative = False

    show_labels = kind == "donut" or (kind == "bar" and point_count <= 40) or (kind == "line" and len(categories) <= 8)
    if show_labels:
        plot.has_data_labels = True
        labels = plot.data_labels
        labels.font.size = Pt(10)
        labels.font.name = FONT
        labels.font.color.rgb = _rgb(TEXT)
        if kind == "donut":
            labels.show_percentage = True
            labels.show_value = False
            labels.number_format = "0%"
        else:
            labels.show_value = True
            # stacked bars: an empty third section hides the "0" labels of empty segments
            labels.number_format = f"{number_format};-{number_format};" if stacked else number_format
            labels.position = (
                XL_LABEL_POSITION.ABOVE if kind == "line"
                else XL_LABEL_POSITION.CENTER if stacked
                else XL_LABEL_POSITION.OUTSIDE_END
            )
        labels.number_format_is_linked = False

    if kind != "donut":
        category_axis = chart.category_axis
        category_axis.tick_labels.font.size = Pt(10)
        category_axis.format.line.color.rgb = _rgb("D9D9D9")
        category_axis.has_major_gridlines = False
        value_axis = chart.value_axis
        value_axis.tick_labels.font.size = Pt(9)
        value_axis.tick_labels.number_format = number_format
        value_axis.tick_labels.number_format_is_linked = False
        value_axis.format.line.fill.background()
        value_axis.has_major_gridlines = not (show_labels and kind == "bar")
        if value_axis.has_major_gridlines:
            value_axis.major_gridlines.format.line.color.rgb = _rgb("E5E5E5")
        if kind == "bar" and show_labels:
            value_axis.visible = False
        if horizontal:
            category_axis.reverse_order = True


def _chart_slide(prs, deck_title: str, spec: dict[str, Any], number: int) -> None:
    slide = _new_slide(prs)
    _frame(slide, deck_title, spec.get("title", ""), number, spec.get("source"))
    caption = _clean(spec.get("chart_title"))
    if caption:
        _text(slide, 0.5, 1.25, 7.7, 0.3, caption, size=12, bold=True, color=GREY)
    _add_chart(slide, spec, 0.5, 1.65, 7.7, 5.2)
    _insight_boxes_right(slide, spec)


# ---------------------------------------------------------------------------
# Tables
# ---------------------------------------------------------------------------
def _is_numeric_text(text: str) -> bool:
    return bool(text) and bool(_NUMERIC_TEXT.match(text))


def _style_cell(cell, text: str, size: float, bold: bool, color: str, fill: str, align) -> None:
    cell.fill.solid()
    cell.fill.fore_color.rgb = _rgb(fill)
    cell.margin_left = cell.margin_right = Inches(0.08)
    cell.margin_top = cell.margin_bottom = Inches(0.03)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame = cell.text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    _style_run(run, size, bold, color)


def _add_table(slide, columns: list[str], rows: list[list[Any]], x: float, y: float, w: float, h: float) -> None:
    columns = [_clean(c) or " " for c in columns][:8]
    rows = [[_clean(v) for v in list(row)[: len(columns)]] for row in rows][:MAX_TABLE_ROWS]
    rows = [row + [""] * (len(columns) - len(row)) for row in rows]
    if not columns or not rows:
        raise ValueError("the table has no data")

    row_count = len(rows) + 1
    size = 12 if row_count <= 7 else 11 if row_count <= 9 else 10
    row_height = min(0.46, max(0.3, h / row_count))
    graphic = slide.shapes.add_table(row_count, len(columns), Inches(x), Inches(y), Inches(w), Inches(row_height * row_count))
    table = graphic.table
    # "No Style, No Grid": every colour below is set by hand
    graphic._element.xpath(".//a:tableStyleId")[0].text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
    table.first_row = True
    table.horz_banding = False

    weights = []
    for index in range(len(columns)):
        longest = max([len(columns[index])] + [len(row[index]) for row in rows])
        weights.append(min(34, max(7, longest)))
    total = sum(weights)
    for index, weight in enumerate(weights):
        table.columns[index].width = Inches(w * weight / total)
    for row in table.rows:
        row.height = Inches(row_height)

    numeric = [all(_is_numeric_text(row[i]) or not row[i] for row in rows) for i in range(len(columns))]
    for index, name in enumerate(columns):
        align = PP_ALIGN.RIGHT if numeric[index] else PP_ALIGN.LEFT
        _style_cell(table.cell(0, index), name, size, True, "FFFFFF", GREEN, align)
    for r, row in enumerate(rows, start=1):
        fill = GREEN_TINT if r % 2 == 1 else "FFFFFF"
        for index, value in enumerate(row):
            align = PP_ALIGN.RIGHT if numeric[index] else PP_ALIGN.LEFT
            _style_cell(table.cell(r, index), value, size, False, TEXT, fill, align)


def _table_slide(prs, deck_title: str, spec: dict[str, Any], number: int) -> None:
    slide = _new_slide(prs)
    _frame(slide, deck_title, spec.get("title", ""), number, spec.get("source"))
    _add_table(slide, spec.get("columns", []), spec.get("rows", []), 0.5, 1.3, 12.33, 3.7)
    _insight_boxes_bottom(slide, spec)


# ---------------------------------------------------------------------------
# Headline numbers
# ---------------------------------------------------------------------------
def _kpi_slide(prs, deck_title: str, spec: dict[str, Any], number: int) -> None:
    kpis = [k for k in list(spec.get("kpis", []))[:MAX_KPIS] if _clean(k.get("value")) and _clean(k.get("label"))]
    if not kpis:
        raise ValueError("there are no headline numbers")
    slide = _new_slide(prs)
    _frame(slide, deck_title, spec.get("title", ""), number, spec.get("source"))

    columns = 4 if len(kpis) > 3 else len(kpis)
    rows = math.ceil(len(kpis) / columns)
    gap = 0.25
    area_w, area_h = 12.33, 3.7
    tile_w = (area_w - gap * (columns - 1)) / columns
    tile_h = min(1.75, (area_h - gap * (rows - 1)) / rows)
    for index, kpi in enumerate(kpis):
        row, col = divmod(index, columns)
        x = 0.5 + col * (tile_w + gap)
        y = 1.35 + row * (tile_h + gap)
        tile = _box(slide, x, y, tile_w, tile_h, TINTS[index % len(TINTS)], PALETTE[index % len(PALETTE)], radius=0.08)
        frame = tile.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = Inches(0.15)
        frame.margin_top = frame.margin_bottom = Inches(0.08)
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        value = _clean(kpi.get("value"))
        value_size = 30 if len(value) <= 9 else 24 if len(value) <= 13 else 18
        paragraph = frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.CENTER
        run = paragraph.add_run()
        run.text = value
        _style_run(run, value_size, True, TEXT)
        paragraph = frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.CENTER
        run = paragraph.add_run()
        run.text = _clean(kpi.get("label"))[:60]
        _style_run(run, 12, False, GREY)
        if _clean(kpi.get("note")):
            paragraph = frame.add_paragraph()
            paragraph.alignment = PP_ALIGN.CENTER
            run = paragraph.add_run()
            run.text = _clean(kpi.get("note"))[:60]
            _style_run(run, 10, False, LIGHT_GREY)
    _insight_boxes_bottom(slide, spec)


# ---------------------------------------------------------------------------
# Closing recommendations
# ---------------------------------------------------------------------------
def _summary_slide(prs, deck_title: str, spec: dict[str, Any], number: int) -> None:
    items = [
        (_clean(i.get("heading")), _clean(i.get("text")))
        for i in list(spec.get("items", []))[:6]
        if _clean(i.get("text")) or _clean(i.get("heading"))
    ]
    if not items:
        raise ValueError("there are no recommendations")
    slide = _new_slide(prs)
    _frame(slide, deck_title, spec.get("title") or "Key recommendations(Experimental Feature)", number, spec.get("source"))

    columns = 2 if len(items) > 1 else 1
    rows = math.ceil(len(items) / columns)
    gap = 0.25
    box_w = (12.33 - gap * (columns - 1)) / columns
    box_h = min(1.9, (5.6 - gap * (rows - 1)) / rows)
    for index, (heading, text) in enumerate(items):
        row, col = divmod(index, columns)
        shape = _box(slide, 0.5 + col * (box_w + gap), 1.3 + row * (box_h + gap), box_w, box_h, BLUE_TINT, BLUE)
        frame = shape.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = Inches(0.2)
        frame.margin_top = frame.margin_bottom = Inches(0.12)
        frame.vertical_anchor = MSO_ANCHOR.TOP
        size, fitted = _fit([text], box_w, box_h, heading=True)
        paragraph = frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.space_after = Pt(5)
        run = paragraph.add_run()
        run.text = f"{index + 1}. {heading}" if heading else str(index + 1)
        _style_run(run, size + 1, True, TEXT)
        paragraph = frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.LEFT
        run = paragraph.add_run()
        run.text = fitted[0]
        _style_run(run, size, False, TEXT)


# ---------------------------------------------------------------------------
# Cover
# ---------------------------------------------------------------------------
def _cover(prs, title: str, subtitle: str, sources: list[str]) -> None:
    layout = _layout(prs, "Cover")
    if layout is not None:
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = title
        if len(title) > 30:
            for run in slide.shapes.title.text_frame.paragraphs[0].runs:
                run.font.size = Pt(28)
        for placeholder in slide.placeholders:
            if placeholder.placeholder_format.type == PP_PLACEHOLDER.SUBTITLE:
                if subtitle:
                    placeholder.text = subtitle
                else:
                    placeholder._element.getparent().remove(placeholder._element)
        if sources:
            _text(slide, 1.73, 6.45, 9.0, 0.4, "Data source: " + ", ".join(sources)[:140], size=11, color="FFFFFF")
        return
    slide = _new_slide(prs)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb(COVER_BLUE)
    _text(slide, 1.2, 2.3, 10.9, 1.6, title, size=40 if len(title) <= 36 else 30, bold=True, color="FFFFFF", anchor=MSO_ANCHOR.BOTTOM)
    if subtitle:
        _text(slide, 1.2, 4.1, 10.9, 0.6, subtitle, size=20, color="FFFFFF")
    if sources:
        _text(slide, 1.2, 6.6, 10.9, 0.4, "Data source: " + ", ".join(sources)[:160], size=11, color="DCEBFF")


def _closing(prs) -> None:
    """Last slide: 'Thank you' on the template's Closing layout (plain blue slide if no template)."""
    layout = _layout(prs, "Closing")
    if layout is not None:
        slide = prs.slides.add_slide(layout)
        if slide.shapes.title is not None:
            slide.shapes.title.text = "Thank you"
        return
    slide = _new_slide(prs)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb(COVER_BLUE)
    _text(slide, 1.2, 2.9, 10.9, 1.6, "Thank you", size=40, bold=True, color="FFFFFF", anchor=MSO_ANCHOR.MIDDLE)


# ---------------------------------------------------------------------------
# Public entry point
# ---------------------------------------------------------------------------
def build_deck(deck: dict[str, Any]) -> tuple[bytes, list[str]]:
    """Returns (the .pptx file as bytes, warnings about slides that had to be skipped)."""
    warnings: list[str] = []
    if TEMPLATE_PATH.exists():
        prs = Presentation(str(TEMPLATE_PATH))
    else:  # no template file: plain deck
        prs = Presentation()
        prs.slide_width = Inches(SLIDE_W)
        prs.slide_height = Inches(SLIDE_H)
    title = _clean(deck.get("title")) or "Report"
    prs.core_properties.title = title
    prs.core_properties.author = "Analytics Copilot"

    sources = _as_list(deck.get("sources"))
    _cover(prs, title, _clean(deck.get("subtitle")), sources)

    number = 1
    for spec in deck.get("slides", []):
        kind = spec.get("type")
        label = _clean(spec.get("title")) or kind or "slide"
        number += 1
        try:
            if kind in ("bar", "line", "donut"):
                _chart_slide(prs, title, spec, number)
            elif kind == "table":
                _table_slide(prs, title, spec, number)
            elif kind == "kpis":
                _kpi_slide(prs, title, spec, number)
            elif kind == "summary":
                _summary_slide(prs, title, spec, number)
            else:
                raise ValueError(f"unknown slide type '{kind}'")
        except Exception as exc:  # noqa: BLE001 - one bad slide must not lose the whole deck
            if len(prs.slides) >= number:
                _drop_last_slide(prs)
            number -= 1
            warnings.append(f"Skipped '{label}': {exc}")

    _closing(prs)

    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue(), warnings